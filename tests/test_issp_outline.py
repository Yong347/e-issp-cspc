from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import html
from sqlalchemy import create_engine, func
from sqlalchemy.orm import Session

from app.database.models import Base, Department, ExistingInventory, MasterNarrative
from app.admin.router import _build_issp_composition_response
from app.services.docx_generator import _apply_docx_outline_overrides
from app.services.issp_outline import (
    apply_pdf_outline_overrides,
    flatten_outline,
    get_default_issp_outline,
    get_issp_outline,
)


def test_default_outline_has_unique_keys_and_anchors():
    nodes = flatten_outline(get_default_issp_outline("2026-2028", 3))

    assert len(nodes) == len({node["key"] for node in nodes})
    assert len(nodes) == len({node["anchor"] for node in nodes})
    assert {"annex1", "annex1-equipment", "annex1-software", "part4-a-year-3"} <= {
        node["key"] for node in nodes
    }


def test_pdf_hides_renames_replaces_and_reorders_sections():
    outline = flatten_outline(get_default_issp_outline("2026-2028", 3))
    anchor_markup = "".join(
        f'<div id="{node["anchor"]}"><h2>{node["title"]}</h2>'
        f'<p>official-{node["key"]}</p></div>'
        for node in outline
    )
    source = (
        '<!doctype html><html><body><div id="sec-toc"><table><tbody>'
        '<tr><td><a href="#sec-supplemental">SUPPLEMENTAL</a></td></tr>'
        '<tr><td><a href="#supplement-0">Replacement</a></td></tr>'
        f"</tbody></table></div>{anchor_markup}"
        '<section id="sec-supplemental"><h1>Supplemental Sections</h1>'
        '<div data-composition-section-id="replacement-1">'
        '<h2 id="supplement-0">Replacement section</h2><p>replacement-content</p>'
        '</div><div data-composition-section-id="subsection-1">'
        '<h2 id="supplement-1">New Annex subsection</h2><p>subsection-content</p>'
        "</div></section></body></html>"
    )
    official_settings = {
        "part2": {"order": 0},
        "definitions": {"order": 1},
        "annex1": {"visible": False, "replacement_section_id": "replacement-1"},
        "part1": {"title": "Agency Context"},
    }

    rendered = apply_pdf_outline_overrides(
        source,
        "2026-2028",
        3,
        official_settings,
        [
            {"id": "replacement-1", "title": "Replacement section", "children": []},
            {"id": "subsection-1", "title": "New Annex subsection", "parent_official_key": "annex1", "children": []},
        ],
    )
    document = html.fromstring(rendered)
    text = " ".join(document.xpath("//body//text()"))

    assert "Agency Context" in text
    assert "official-annex1" not in text
    assert "replacement-content" in text
    assert text.index("Replacement section") < text.index("New Annex subsection")
    assert "subsection-content" in text
    assert text.index("PART II. CURRENT ICT ASSESSMENT") < text.index("Definition of Terms")
    assert not document.xpath("//*[@id='supplement-0']")


def test_docx_hides_renames_replaces_and_reorders_sections():
    document = Document()
    for text in (
        "DEFINITION OF TERMS",
        "PART I. AGENCY PROFILE & STRATEGIC CONTEXT",
        "PART II. CURRENT ICT ASSESSMENT",
        "PART III. PROPOSED ICT STRATEGY",
        "PART IV. RESOURCE REQUIREMENTS & SUMMARY OF INVESTMENTS",
        "A.1. [YEAR #1] (FY 2026)",
        "ANNEX 1: EXISTING INFORMATION & COMMUNICATIONS TECHNOLOGY (ICT) ASSET INVENTORY",
        "original inventory content",
        "ANNEX 2: DISASTER RECOVERY & BUSINESS CONTINUITY PLAN",
        "Supplemental Sections",
    ):
        document.add_paragraph(text)

    for section_index, section_title, content in (
        (0, "Replacement section", "replacement inventory content"),
        (1, "New Annex subsection", "nested replacement content"),
    ):
        start = document.add_paragraph()
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(section_index * 2 + 1))
        bookmark_start.set(qn("w:name"), f"issp_custom_section_start_{section_index}")
        start._p.append(bookmark_start)
        document.add_paragraph(section_title)
        document.add_paragraph(content)
        end = document.add_paragraph()
        bookmark_end = OxmlElement("w:bookmarkStart")
        bookmark_end.set(qn("w:id"), str(section_index * 2 + 2))
        bookmark_end.set(qn("w:name"), f"issp_custom_section_end_{section_index}")
        end._p.append(bookmark_end)

    _apply_docx_outline_overrides(
        document,
        "2026-2028",
        3,
        {
            "part2": {"order": 0},
            "definitions": {"order": 1},
            "part1": {"title": "Agency Context"},
            "annex1": {"visible": False, "replacement_section_id": "replacement-1"},
        },
        [
            {"id": "replacement-1", "title": "Replacement section", "children": []},
            {"id": "subsection-1", "title": "New Annex subsection", "parent_official_key": "annex1", "children": []},
        ],
    )
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert paragraphs.index("PART II. CURRENT ICT ASSESSMENT") < paragraphs.index("DEFINITION OF TERMS")
    assert "Agency Context" in paragraphs
    assert "original inventory content" not in paragraphs
    assert "replacement inventory content" in paragraphs
    assert "nested replacement content" in paragraphs
    assert paragraphs.index("Replacement section") < paragraphs.index("New Annex subsection")
    assert paragraphs.index("nested replacement content") < paragraphs.index(
        "ANNEX 2: DISASTER RECOVERY & BUSINESS CONTINUITY PLAN"
    )


def test_outline_settings_are_cycle_scoped_and_do_not_delete_source_rows():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)

    with Session(engine) as db:
        db.add(Department(id=1, name="Outline Test Department"))
        db.add_all([
            MasterNarrative(issp_period="2026-2028", issp_period_from=2026, issp_period_to=2028),
            MasterNarrative(issp_period="2029-2031", issp_period_from=2029, issp_period_to=2031),
        ])
        db.flush()
        db.add(ExistingInventory(
            department_id=1,
            issp_period="2026-2028",
            item_name="Test laptop",
            is_software=False,
        ))
        db.flush()

        configured = get_issp_outline(db, "2026-2028", 3, {"annex1": {"visible": False}})
        annex1 = next(section for section in configured if section["key"] == "annex1")
        response = _build_issp_composition_response(
            db,
            db.query(MasterNarrative).filter(MasterNarrative.issp_period == "2026-2028").one(),
            {"version": 1, "sections": []},
            None,
        )
        other_cycle = get_issp_outline(db, "2029-2031", 3)
        other_annex1 = next(section for section in other_cycle if section["key"] == "annex1")

        assert annex1["visible"] is False
        assert annex1["record_count"] == 1
        assert annex1["children"][0]["records"][0]["item_name"] == "Test laptop"
        api_annex1 = next(section for section in response.outline if section.key == "annex1")
        assert response.version == 1
        assert api_annex1.children[0].records[0]["item_name"] == "Test laptop"
        assert other_annex1["record_count"] == 0
        assert other_annex1["children"][0]["records"] == []
        assert db.query(func.count(ExistingInventory.id)).scalar() == 1